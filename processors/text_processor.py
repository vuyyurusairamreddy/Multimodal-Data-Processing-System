import PyPDF2
import pdfplumber
import docx
from pptx import Presentation
from config import SUPPORTED_TEXT_FORMATS

def process_pdf(file_path):
    text = ""
    
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
    except:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                text += page.extract_text() + "\n"
    
    return text.strip()

def process_docx(file_path):
    doc = docx.Document(file_path)
    text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
    return text.strip()

def process_pptx(file_path):
    prs = Presentation(file_path)
    text = ""
    
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    
    return text.strip()

def process_txt(file_path):
    with open(file_path, 'r', encoding='utf-8') as file:
        text = file.read()
    return text.strip()

def process_text_file(file_path, file_extension):
    if file_extension == '.pdf':
        return process_pdf(file_path)
    elif file_extension == '.docx':
        return process_docx(file_path)
    elif file_extension == '.pptx':
        return process_pptx(file_path)
    elif file_extension in ['.txt', '.md']:
        return process_txt(file_path)
    else:
        return ""
